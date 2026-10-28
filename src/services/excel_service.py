"""Servicio de generación de planillas Excel nativas (.xlsx) para horas laborales.

Construye libros de cálculo profesionales con fórmulas activas, estilos de celda,
alineaciones ergonómicas y el cuadro resumen de liquidación e IVA estandarizado.
"""

from __future__ import annotations

import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from logger_config import setup_logger

logger = setup_logger("excel_service")

BASE_DIR = Path(__file__).resolve().parents[1]
REPORTES_DIR = BASE_DIR / "reportes"


@dataclass
class DatosPlanillaHoras:
    """Contenedor de datos estructurados para compilar la planilla Excel."""

    periodo: str
    nombre_titular: str
    encabezados: List[str]
    filas: List[List[Any]]
    monto_por_hora: int
    saldo_anterior: int = 0
    descuento: int = 0
    nombre_archivo: Optional[str] = None


class ExcelReportService:
    """Generador de libros Excel (.xlsx) para rendición de asistencia y horas."""

    COLOR_TITULO = "D9D9D9"
    COLOR_ENCABEZADO = "B0B0B0"
    COLOR_DESTACADO = "A6A6A6"
    FORMATO_MONEDA = '"$"#,##0'
    FORMATO_HORAS = "0.00"

    def __init__(self, output_dir: Path | str | None = None) -> None:
        self.output_dir = Path(output_dir) if output_dir else REPORTES_DIR

    def generar_reporte_horas(self, datos: DatosPlanillaHoras | Any) -> str:
        """Genera la planilla Excel y retorna la ruta absoluta del archivo generado."""
        if not isinstance(datos, DatosPlanillaHoras):
            resumen = getattr(datos, "resumen_calculado", {}) or {}
            datos = DatosPlanillaHoras(
                periodo=getattr(datos, "subtitulo", "")
                or getattr(datos, "periodo", ""),
                nombre_titular=getattr(datos, "nombre_titular", "KEVIN PERALTA"),
                encabezados=getattr(datos, "encabezados", []),
                filas=getattr(datos, "filas", []),
                monto_por_hora=resumen.get("hora", 14634),
                saldo_anterior=resumen.get("saldo_anterior", 0),
                descuento=resumen.get("descuento", 0),
                nombre_archivo=getattr(datos, "nombre_archivo", None),
            )
        ruta, msg = self.generar_reporte_horas_excel(datos, output_dir=self.output_dir)
        if not ruta:
            raise RuntimeError(f"Error generando reporte Excel: {msg}")
        return ruta

    @classmethod
    def generar_reporte_horas_excel(
        cls, datos: DatosPlanillaHoras, output_dir: Path | str | None = None
    ) -> tuple[Optional[str], str]:
        """Genera un archivo .xlsx estilizado con fórmulas y cuadro resumen.

        Retorna (ruta_archivo_xlsx, mensaje_estado).
        """
        try:
            target_dir = Path(output_dir) if output_dir else REPORTES_DIR
            target_dir.mkdir(parents=True, exist_ok=True)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = datos.periodo[:31]  # Límite de caracteres en Excel

            # ── Estilos reutilizables ──────────────────────────────────────────
            borde_fino = Side(border_style="thin", color="000000")
            borde_doble = Side(border_style="double", color="000000")
            border_celda = Border(
                left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino
            )
            border_destacado = Border(
                left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_doble
            )

            fill_titulo = PatternFill(
                start_color=cls.COLOR_TITULO,
                end_color=cls.COLOR_TITULO,
                fill_type="solid",
            )
            fill_encabezado = PatternFill(
                start_color=cls.COLOR_ENCABEZADO,
                end_color=cls.COLOR_ENCABEZADO,
                fill_type="solid",
            )
            fill_destacado = PatternFill(
                start_color=cls.COLOR_DESTACADO,
                end_color=cls.COLOR_DESTACADO,
                fill_type="solid",
            )

            font_titulo = Font(name="Calibri", size=12, bold=True)
            font_encabezado = Font(name="Calibri", size=10, bold=True)
            font_datos = Font(name="Calibri", size=10)
            font_destacado = Font(name="Calibri", size=11, bold=True)

            align_center = Alignment(horizontal="center", vertical="center")
            align_left = Alignment(horizontal="left", vertical="center")
            align_right = Alignment(horizontal="right", vertical="center")

            # ── 1. Título General (Fila 1 fusionada A1:G1) ─────────────────────
            num_cols = max(len(datos.encabezados), 7)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
            titulo_texto = (
                f"REGISTRO DE HORAS DEL MES DE {datos.periodo.upper()} - "
                f"{datos.nombre_titular.upper()}"
            )
            celda_titulo = ws.cell(row=1, column=1, value=titulo_texto)
            celda_titulo.font = font_titulo
            celda_titulo.alignment = align_center
            celda_titulo.fill = fill_titulo
            ws.row_dimensions[1].height = 24

            for col in range(1, num_cols + 1):
                c = ws.cell(row=1, column=col)
                c.border = border_celda
                c.fill = fill_titulo

            # ── 2. Encabezados de Columnas (Fila 2) ─────────────────────────────
            ws.row_dimensions[2].height = 20
            for col_idx, encabezado in enumerate(datos.encabezados, start=1):
                c = ws.cell(row=2, column=col_idx, value=encabezado)
                c.font = font_encabezado
                c.alignment = align_center
                c.fill = fill_encabezado
                c.border = border_celda

            # ── 3. Filas de Datos (Fila 3 en adelante) ──────────────────────────
            fila_actual = 3
            indice_col_horas = None
            for idx, h in enumerate(datos.encabezados, start=1):
                if (
                    "hora" in h.lower()
                    and "salida" not in h.lower()
                    and "entrada" not in h.lower()
                ):
                    indice_col_horas = idx
                    break
            if indice_col_horas is None:
                indice_col_horas = 7  # Por defecto columna G

            letra_col_horas = get_column_letter(indice_col_horas)

            for f_data in datos.filas:
                ws.row_dimensions[fila_actual].height = 18
                for col_idx in range(1, num_cols + 1):
                    val = f_data[col_idx - 1] if (col_idx - 1) < len(f_data) else ""
                    c = ws.cell(row=fila_actual, column=col_idx)

                    # Si es la columna de horas y es convertible a float
                    if col_idx == indice_col_horas and val != "":
                        try:
                            # Reemplazar comas si vienen con formato latino
                            val_clean = str(val).replace(",", ".").strip()
                            num_horas = float(val_clean)
                            c.value = num_horas
                            c.number_format = cls.FORMATO_HORAS
                            c.alignment = align_right
                        except ValueError:
                            c.value = val
                            c.alignment = align_center
                    else:
                        c.value = val
                        c.alignment = align_center

                    c.font = font_datos
                    c.border = border_celda
                fila_actual += 1

            ultima_fila_datos = max(3, fila_actual - 1)

            # ── 4. Cuadro Resumen Inferior ─────────────────────────────────────
            # Dejamos una fila en blanco
            fila_resumen = fila_actual + 1

            # Cabecera de Cuadro Resumen
            c_res = ws.cell(row=fila_resumen, column=1, value="Resumen")
            c_res.font = font_encabezado
            c_res.alignment = align_center
            c_res.fill = fill_encabezado
            c_res.border = border_celda

            c_cant = ws.cell(row=fila_resumen, column=2, value="Cantidades")
            c_cant.font = font_encabezado
            c_cant.alignment = align_center
            c_cant.fill = fill_encabezado
            c_cant.border = border_celda

            # Fila: mes (Total horas trabajadas)
            fila_mes = fila_resumen + 1
            ws.cell(row=fila_mes, column=1, value="mes").alignment = align_left
            ws.cell(row=fila_mes, column=1).font = font_datos
            ws.cell(row=fila_mes, column=1).border = border_celda

            c_mes_val = ws.cell(
                row=fila_mes,
                column=2,
                value=f"=SUM({letra_col_horas}3:{letra_col_horas}{ultima_fila_datos})",
            )
            c_mes_val.font = font_datos
            c_mes_val.alignment = align_right
            c_mes_val.number_format = "0.0"
            c_mes_val.border = border_celda

            # Fila: hora (Monto tarifa por hora)
            fila_hora = fila_mes + 1
            ws.cell(row=fila_hora, column=1, value="hora").alignment = align_left
            ws.cell(row=fila_hora, column=1).font = font_datos
            ws.cell(row=fila_hora, column=1).border = border_celda

            c_hora_val = ws.cell(row=fila_hora, column=2, value=datos.monto_por_hora)
            c_hora_val.font = font_datos
            c_hora_val.alignment = align_right
            c_hora_val.number_format = cls.FORMATO_MONEDA
            c_hora_val.border = border_celda

            # Fila: por Horas (Suma horas * Tarifa)
            fila_por_horas = fila_hora + 1
            ws.cell(row=fila_por_horas, column=1, value="por Horas").alignment = (
                align_left
            )
            ws.cell(row=fila_por_horas, column=1).font = font_datos
            ws.cell(row=fila_por_horas, column=1).border = border_celda

            c_por_h_val = ws.cell(
                row=fila_por_horas,
                column=2,
                value=f"=ROUND(B{fila_mes}*B{fila_hora}, 0)",
            )
            c_por_h_val.font = font_datos
            c_por_h_val.alignment = align_right
            c_por_h_val.number_format = cls.FORMATO_MONEDA
            c_por_h_val.border = border_celda

            # Fila: IVA 10%
            fila_iva = fila_por_horas + 1
            ws.cell(row=fila_iva, column=1, value="IVA 10%").alignment = align_left
            ws.cell(row=fila_iva, column=1).font = font_datos
            ws.cell(row=fila_iva, column=1).border = border_celda

            c_iva_val = ws.cell(
                row=fila_iva,
                column=2,
                value=f"=ROUND(B{fila_por_horas}*0.1, 0)",
            )
            c_iva_val.font = font_datos
            c_iva_val.alignment = align_right
            c_iva_val.number_format = cls.FORMATO_MONEDA
            c_iva_val.border = border_celda

            # Fila: INCLUIDO
            fila_incluido = fila_iva + 1
            ws.cell(row=fila_incluido, column=1, value="INCLUIDO").alignment = (
                align_left
            )
            ws.cell(row=fila_incluido, column=1).font = font_datos
            ws.cell(row=fila_incluido, column=1).border = border_celda

            c_inc_val = ws.cell(
                row=fila_incluido,
                column=2,
                value=f"=B{fila_por_horas}+B{fila_iva}",
            )
            c_inc_val.font = font_datos
            c_inc_val.alignment = align_right
            c_inc_val.number_format = cls.FORMATO_MONEDA
            c_inc_val.border = border_celda

            # Fila Opcional: SALDO ANTERIOR o Descuento (Solo si > 0)
            fila_ajuste = None
            es_descuento = False
            if datos.saldo_anterior > 0:
                fila_ajuste = fila_incluido + 1
                c_lbl = ws.cell(row=fila_ajuste, column=1, value="SALDO ANTERIOR")
                c_lbl.font = font_datos
                c_lbl.alignment = align_left
                c_lbl.border = border_celda

                c_val = ws.cell(
                    row=fila_ajuste, column=2, value=int(datos.saldo_anterior)
                )
                c_val.font = font_datos
                c_val.alignment = align_right
                c_val.number_format = cls.FORMATO_MONEDA
                c_val.border = border_celda
            elif datos.descuento > 0:
                fila_ajuste = fila_incluido + 1
                es_descuento = True
                c_lbl = ws.cell(row=fila_ajuste, column=1, value="Descuento")
                c_lbl.font = font_datos
                c_lbl.alignment = align_left
                c_lbl.border = border_celda

                c_val = ws.cell(row=fila_ajuste, column=2, value=-int(datos.descuento))
                c_val.font = font_datos
                c_val.alignment = align_right
                c_val.number_format = cls.FORMATO_MONEDA
                c_val.border = border_celda

            # Fila Final: FACTURA
            fila_factura = (fila_ajuste + 1) if fila_ajuste else (fila_incluido + 1)
            c_fac_lbl = ws.cell(row=fila_factura, column=1, value="FACTURA")
            c_fac_lbl.font = font_destacado
            c_fac_lbl.alignment = align_left
            c_fac_lbl.fill = fill_destacado
            c_fac_lbl.border = border_destacado

            if fila_ajuste:
                if es_descuento:
                    formula_final = f"=B{fila_incluido}+B{fila_ajuste}"  # B{fila_ajuste} ya es negativo
                else:
                    formula_final = f"=B{fila_incluido}+B{fila_ajuste}"
            else:
                formula_final = f"=B{fila_incluido}"

            c_fac_val = ws.cell(row=fila_factura, column=2, value=formula_final)
            c_fac_val.font = font_destacado
            c_fac_val.alignment = align_right
            c_fac_val.fill = fill_destacado
            c_fac_val.number_format = cls.FORMATO_MONEDA
            c_fac_val.border = border_destacado

            # ── 5. Ajuste Ergonómico de Columnas ───────────────────────────────
            anchos_predeterminados = {1: 14, 2: 14, 3: 15, 4: 15, 5: 15, 6: 15, 7: 12}
            for col_num in range(1, num_cols + 1):
                col_letter = get_column_letter(col_num)
                ancho = anchos_predeterminados.get(col_num, 14)
                ws.column_dimensions[col_letter].width = ancho

            # ── 6. Sanitizar y Guardar Archivo ─────────────────────────────────
            nombre_base = (
                datos.nombre_archivo
                or f"reporte_horas_{datos.periodo.replace(' ', '_')}"
            )
            nombre_sano = re.sub(r"[^\w\-\.]", "_", nombre_base)
            if not nombre_sano.endswith(".xlsx"):
                nombre_sano += ".xlsx"

            sufijo_id = uuid.uuid4().hex[:6]
            stem = Path(nombre_sano).stem
            nombre_final = f"{stem}_{sufijo_id}.xlsx"
            ruta_salida = target_dir / nombre_final

            wb.save(ruta_salida)
            logger.info(f"✅ Planilla Excel generada exitosamente en {ruta_salida}")
            return str(ruta_salida), f"Planilla Excel generada: {nombre_final}"

        except Exception as e:
            logger.error(f"❌ Error al generar planilla Excel: {e}")
            return None, f"Error generando Excel: {str(e)}"
